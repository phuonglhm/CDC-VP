#!/usr/bin/env python3
"""Generate the FX1 ISP CSR artefacts from the pinned CSR workbook (FTEL_IP_ISP_CSR_v1.07.xlsm).

The spreadsheet is the ABI source (decision DEC-09: CSR wins over the HAS).
It is confidential and is NOT stored in this repository; pass its path on the
command line. The macro-enabled .xlsm is read as data only (openpyxl; no macro
is executed). The generated files carry the ABI only -- names, offsets, bit
positions, access types, reset values, enum names and the worksheet row of
each item -- and never the spreadsheet's descriptive text. The generated files are kept in the source tree so that the C++
build never needs Python or the spreadsheet:

  docs/csr/fx1_isp_csr_schema.json   machine-readable schema, cell traceability
  include/fx1_isp/fx1_isp_csr.h      C/C++ register header for FW/SW
  src/registers/csr_table_gen.cpp    register table consumed by the model

Every deviation from the literal spreadsheet content is an explicit entry in
OVERRIDES below, carries a decision/issue ID, and is asserted against the
spreadsheet text so that a new revision of the map fails loudly instead of
silently inheriting a stale override.

Usage:
  gen_csr.py --xlsx <path/to/FTEL_IP_ISP_CSR_v1.07.xlsm> [--out <isp dir>] [--check]
"""

import argparse
import hashlib
import json
import os
import re
import sys
import tempfile

try:
    import openpyxl
except ImportError:  # pragma: no cover - environment problem, not a bug
    sys.exit("gen_csr.py: openpyxl is required (pip install openpyxl)")

GENERATOR_VERSION = "2.0"
# CSR revision v1.07 (ABI 0x00010007, MAS v1.07 of 2026-10-18). Previous pins:
# bda077b6... (07/10, 254 registers, 336 fields; docs/ISP_CSR_REVISION_0710.md)
# and 0980b61d... (226 registers). See docs/ISP_CSR_REVISION_V107.md.
EXPECTED_SHA256 = "444e72f8d7fd3f35d83a7ce47d455b87dc6df6cda833b11302c6a45fb8b721fa"
EXPECTED_COUNTS = {"registers": 573, "fields": 819,
                   "RW": 464, "RO": 288, "W1C": 35, "W1S": 11, "W1SC": 2, "W1P": 19}
# MAS v1.07 Table 8-178. W1P: one-clock command, reads zero (new in v1.07).
ACCESS_TYPES = ("RW", "RO", "W1C", "W1S", "W1SC", "W1P")
APERTURE_BYTES = 0x10000

# Overrides of the literal spreadsheet content. Each one is checked against the
# text it overrides (`expect_*`), so it cannot outlive the defect it corrects.
#
# Kinds:
#   assert_enum  no change; asserts the spreadsheet enum (a resolved decision
#                that must not silently regress in a later revision)
#   enum_names   the description gives values in prose; name them for the header
#   bit_names    the description gives bit meanings without names; name them
#   tied_zero    the bits are reserved and read zero
#
# CSR 07/10: DEC-11 is now the spreadsheet's own enum order (assert only);
# SPEC-05 is superseded -- DMA_IRQ bit 0 is CFG_ERROR, no longer reserved.
# CSR v1.07: buffer-ownership bits at or above the build buffer count (4,
# PARA_DMA_NUM_BUFFERS) read zero and ignore writes (V107-1).
BUFFER_FIELDS = [("IDMA_BUF_VALID", "idma_buf_valid", "7:4 unused"),
                 ("ODMA_BUF_FREE", "odma_buf_free", "bits 7:4 are unused"),
                 ("ODMA_BUF_DONE", "odma_buf_done", "bits 7:4 unused"),
                 ("CH1_ODMA_BUF_FREE", "odma_buf_free", "bits 7:4 are unused"),
                 ("CH1_ODMA_BUF_DONE", "odma_buf_done", "bits 7:4 unused"),
                 ("CH2_ODMA_BUF_FREE", "odma_buf_free", "bits 7:4 are unused"),
                 ("CH2_ODMA_BUF_DONE", "odma_buf_done", "bits 7:4 unused")]
# MAS v1.07 Table 6-122 (codes 0xF and 0x16..0x1F reserved).
RESIZER_PRESETS = {0: "FULLRES", 1: "RES_3840X2160", 2: "RES_2560X1440", 3: "RES_2880X1620", 4: "RES_2304X1296",
                   5: "RES_1920X1080", 6: "RES_1280X720", 7: "RES_960X540", 8: "RES_640X360", 9: "RES_2592X1944",
                   10: "RES_2048X1536", 11: "RES_1600X1200", 12: "RES_1280X960", 13: "RES_800X600",
                   14: "RES_640X480", 16: "RES_640X640", 17: "RES_640X384", 18: "RES_640X352", 19: "RES_512X512",
                   20: "RES_416X416", 21: "RES_320X320"}
OVERRIDES = [
    {
        "id": "DEC-11",
        "register": "COMMON_BAYER", "field": "pattern", "kind": "assert_enum",
        "reason": "The [V,H] encoding 0=RGGB 1=GRBG 2=GBRG 3=BGGR adopted by DEC-11 is the "
                  "spreadsheet's own enum since the 07/10 revision; asserted so it cannot regress.",
        "expect_enum": {0: "RGGB", 1: "GRBG", 2: "GBRG", 3: "BGGR"},
    },
    {
        "id": "CSR-DMA-ERR-5",
        "register": "DMA_ERR", "field": "err", "kind": "tied_zero",
        "reason": "Cause list marks bit 5 as reserved, reads zero.",
        "expect_text": "Bit 5 reserved, reads zero.",
        "bits": [5],
    },
    {
        "id": "VP-NAMES-CSC",
        "register": "CSC_CTRL", "field": "std", "kind": "enum_names",
        "reason": "The values are given in prose; keep the header names BT601/BT709.",
        "expect_text": "0 BT.601 (reset), 1 BT.709.",
        "enum": {0: "BT601", 1: "BT709"},
    },
    {
        "id": "VP-NAMES-INPUT-DEPTH",
        "register": "COMMON_INPUT_FORMAT", "field": "DEPTH", "kind": "enum_names",
        "reason": "The codes are given in prose; name the two legal native depths (2/3 invalid).",
        "expect_text": "0 RAW12 unchanged; 1 RAW10 shifted left two; 2/3 invalid.",
        "enum": {0: "RAW12", 1: "RAW10"},
    },
    {
        "id": "VP-NAMES-ODMA-BUF",
        "register": "ODMA_BUF_FREE", "field": "odma_buf_free", "kind": "bit_names",
        "reason": "Bits 3:0 identify the four buffers without names; keep BUFFER_0..3.",
        "expect_text": "Bits 3:0 identify the four buffers",
        "bits": {0: "BUFFER_0", 1: "BUFFER_1", 2: "BUFFER_2", 3: "BUFFER_3"},
    },
    {
        "id": "VP-NAMES-ODMA-BUF",
        "register": "ODMA_BUF_DONE", "field": "odma_buf_done", "kind": "bit_names",
        "reason": "Bits 3:0 identify the buffers without names; keep BUFFER_0..3.",
        "expect_text": "Bits 3:0 identify buffers",
        "bits": {0: "BUFFER_0", 1: "BUFFER_1", 2: "BUFFER_2", 3: "BUFFER_3"},
    },
] + [
    {"id": "VP-NAMES-ODMA-BUF", "register": reg, "field": fld, "kind": "bit_names",
     "reason": "Bits 3:0 identify the buffers without names; BUFFER_0..3 as on CH0.",
     "expect_text": "Bits 3:0 identify " + ("the four buffers" if fld == "odma_buf_free" else "buffers"),
     "bits": {0: "BUFFER_0", 1: "BUFFER_1", 2: "BUFFER_2", 3: "BUFFER_3"}}
    for reg, fld in (("CH1_ODMA_BUF_FREE", "odma_buf_free"), ("CH1_ODMA_BUF_DONE", "odma_buf_done"),
                     ("CH2_ODMA_BUF_FREE", "odma_buf_free"), ("CH2_ODMA_BUF_DONE", "odma_buf_done"))
] + [
    {"id": "VP-NAMES-RESIZER", "register": reg, "field": "scale", "kind": "enum_names",
     "reason": "The preset codes are given in prose; name them as in MAS v1.07 Table 6-122 "
               "(1..14 unchanged from 07/10, 16..21 the AI presets; CH2 only its legal codes).",
     "expect_text": text, "enum": enum}
    for reg, text, enum in (
        ("RESIZER_CTRL", "16=640x640, 17=640x384, 18=640x352, 19=512x512, 20=416x416, 21=320x320", RESIZER_PRESETS),
        ("CH1_RESIZER_CTRL", "16=640x640, 17=640x384, 18=640x352, 19=512x512, 20=416x416, 21=320x320",
         RESIZER_PRESETS),
        ("CH2_RESIZER_CTRL", "8=640x360; 14=640x480; 16=640x640; 17=640x384; 18=640x352; 19=512x512; "
                             "20=416x416; 21=320x320",
         {k: v for k, v in RESIZER_PRESETS.items() if k in (0, 8, 14, 16, 17, 18, 19, 20, 21)}))
] + [
    {"id": "CSR107-BUFCOUNT", "register": reg, "field": fld, "kind": "tied_zero",
     "reason": "Bits at or above the build buffer count (4) read zero and ignore writes.",
     "expect_text": text, "bits": [4, 5, 6, 7]}
    for reg, fld, text in BUFFER_FIELDS
]

# Header names that a revision renamed. The old macro stays as an alias of the
# new one so that code written against the previous map keeps compiling; each
# entry is asserted (old name gone, new name present) like an override.
COMPAT_ALIASES = [
    {"id": "CSR0710-RENAME", "register": "DMA_ERR", "field": "err", "old": "ALIGN_OR_GEOMETRY", "new": "ALIGN"},
    {"id": "CSR0710-RENAME", "register": "DMA_ERR", "field": "err", "old": "IDMA_AXI", "new": "AXI_RD"},
    {"id": "CSR0710-RENAME", "register": "DMA_ERR", "field": "err", "old": "ODMA_AXI", "new": "AXI_WR"},
]

ENUM_VALUE_RE = re.compile(r"^\s*0x([0-9A-Fa-f]+)\s*=\s*([A-Za-z][A-Za-z0-9_]*)")
ENUM_BIT_RE = re.compile(r"^\s*bit\s+(\d+)\s*=\s*([A-Za-z][A-Za-z0-9_]*)")
# 07/10 style: "Bit 0 CFG_ERROR (...). Bit 1 IDMA_START." or "Bit 3 AXI_RD: ...",
# possibly several per line; the name is an upper-case identifier.
BIT_INLINE_RE = re.compile(r"\bBit\s+(\d+)\s+([A-Z][A-Z0-9_]+)\b")
NOT_IMPL_RE = re.compile(r"not implemented|no hardware|no processing effect|compatibility field"
                         r"|reserved for a future", re.IGNORECASE)


def fail(msg):
    sys.exit("gen_csr.py: ERROR: " + msg)


def sha256_of(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def parse_int(text, what):
    try:
        return int(str(text).strip(), 0)
    except ValueError:
        fail("cannot parse %s %r" % (what, text))


def parse_bits(text, what):
    m = re.fullmatch(r"\[(\d+)(?::(\d+))?\]", str(text).strip())
    if not m:
        fail("cannot parse bit range %r of %s" % (text, what))
    msb = int(m.group(1))
    lsb = int(m.group(2)) if m.group(2) is not None else msb
    if lsb > msb:
        fail("inverted bit range %r of %s" % (text, what))
    return msb, lsb


def read_sheet(ws):
    """Return the register list of one worksheet, with source row numbers."""
    header = [c.value for c in ws[1]]
    want = ["Address", "Register Name", "Field Name", "Bit Range", "Width",
            "Type", "Reset Value", "Description"]
    if header[:8] != want:
        fail("unexpected header in sheet %r: %r" % (ws.title, header))
    regs = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        addr, rname, fname, bits, width, acc, rst, desc = row[:8]
        if all(v is None for v in row[:8]):
            continue
        desc = (desc or "").strip()
        if rname is not None:
            regs.append({"offset": parse_int(addr, "address"), "name": str(rname).strip(),
                         "description": desc, "row": row_idx, "fields": []})
            continue
        if not regs:
            fail("field row %d before any register" % row_idx)
        reg = regs[-1]
        what = "%s.%s" % (reg["name"], fname)
        msb, lsb = parse_bits(bits, what)
        if width is None or int(width) != msb - lsb + 1:
            fail("width %r disagrees with bit range %r for %s" % (width, bits, what))
        acc = str(acc).strip()
        if acc not in ACCESS_TYPES:
            fail("unknown access type %r for %s" % (acc, what))
        enums, bit_names = {}, {}
        for line in desc.splitlines():
            m = ENUM_VALUE_RE.match(line)
            if m:
                enums[int(m.group(1), 16)] = m.group(2)
                continue
            m = ENUM_BIT_RE.match(line)
            if m:
                bit_names[int(m.group(1))] = m.group(2)
                continue
            for b, name in BIT_INLINE_RE.findall(line):
                bit_names.setdefault(int(b), name)
        reg["fields"].append({
            "name": str(fname).strip(), "msb": msb, "lsb": lsb, "width": msb - lsb + 1,
            "access": acc, "reset": parse_int(rst, "reset of " + what),
            "description": desc, "row": row_idx, "enums": enums, "bits": bit_names,
            "implemented": not (NOT_IMPL_RE.search(desc) or
                                "NOT IMPLEMENTED" in reg["description"]),
            "tied_zero_bits": [],
        })
    return regs


def normalise(regs):
    return [(r["offset"], r["name"], r["description"],
             [(f["name"], f["msb"], f["lsb"], f["access"], f["reset"], f["description"])
              for f in r["fields"]]) for r in regs]


def apply_overrides(regs):
    by_name = {r["name"]: r for r in regs}
    applied = []
    for ov in OVERRIDES:
        reg = by_name.get(ov["register"])
        if reg is None:
            fail("override %s: register %s missing" % (ov["id"], ov["register"]))
        fld = next((f for f in reg["fields"] if f["name"] == ov["field"]), None)
        if fld is None:
            fail("override %s: field %s.%s missing" % (ov["id"], ov["register"], ov["field"]))
        if ov["kind"] == "assert_enum":
            if fld["enums"] != ov["expect_enum"]:
                fail("override %s: %s.%s enum is %r, expected %r -- re-review the map"
                     % (ov["id"], ov["register"], ov["field"], fld["enums"], ov["expect_enum"]))
        elif ov["kind"] in ("enum_names", "bit_names"):
            if ov["expect_text"] not in fld["description"]:
                fail("override %s: text %r not found in %s.%s -- re-review the map"
                     % (ov["id"], ov["expect_text"], ov["register"], ov["field"]))
            have = fld["enums"] if ov["kind"] == "enum_names" else fld["bits"]
            if have:
                fail("override %s: %s.%s now names its values (%r) -- drop the override"
                     % (ov["id"], ov["register"], ov["field"], have))
            if ov["kind"] == "enum_names":
                fld["enums"] = dict(ov["enum"])
            else:
                fld["bits"] = dict(ov["bits"])
        elif ov["kind"] == "tied_zero":
            if ov["expect_text"] not in fld["description"]:
                fail("override %s: text %r not found in %s.%s -- re-review the map"
                     % (ov["id"], ov["expect_text"], ov["register"], ov["field"]))
            for b in ov["bits"]:
                if not fld["lsb"] <= b <= fld["msb"]:
                    fail("override %s: bit %d outside field" % (ov["id"], b))
            fld["tied_zero_bits"] = sorted(set(fld["tied_zero_bits"]) | set(ov["bits"]))
        else:
            fail("unknown override kind %r" % ov["kind"])
        applied.append({k: v for k, v in ov.items() if not k.startswith("expect_")})
    for al in COMPAT_ALIASES:
        fld = next(f for f in by_name[al["register"]]["fields"] if f["name"] == al["field"])
        names = set(fld["bits"].values()) | set(fld["enums"].values())
        if al["new"] not in names or al["old"] in names:
            fail("alias %s: %s.%s names %r; expected %s present and %s gone -- re-review"
                 % (al["id"], al["register"], al["field"], sorted(names), al["new"], al["old"]))
    return applied


def field_mask(f):
    mask = ((1 << f["width"]) - 1) << f["lsb"]
    for b in f["tied_zero_bits"]:
        mask &= ~(1 << b)
    return mask


def validate(regs):
    names, offsets = set(), set()
    counts = {a: 0 for a in ACCESS_TYPES}
    prev = -1
    for r in regs:
        off = r["offset"]
        if off % 4 or off >= APERTURE_BYTES:
            fail("register %s offset 0x%X misaligned or outside aperture" % (r["name"], off))
        if off <= prev:
            fail("register %s offset 0x%X not ascending" % (r["name"], off))
        prev = off
        if r["name"] in names or off in offsets:
            fail("duplicate register name/offset %s 0x%X" % (r["name"], off))
        names.add(r["name"])
        offsets.add(off)
        if not r["fields"]:
            fail("register %s has no fields" % r["name"])
        used, fnames = 0, set()
        for f in r["fields"]:
            if f["msb"] > 31:
                fail("%s.%s exceeds 32 bits" % (r["name"], f["name"]))
            if f["name"] in fnames:
                fail("duplicate field %s.%s" % (r["name"], f["name"]))
            fnames.add(f["name"])
            full = ((1 << f["width"]) - 1) << f["lsb"]
            if used & full:
                fail("overlapping field %s.%s" % (r["name"], f["name"]))
            used |= full
            if f["reset"] >> f["width"]:
                fail("reset 0x%X does not fit %s.%s" % (f["reset"], r["name"], f["name"]))
            counts[f["access"]] += 1
    got = dict(counts, registers=len(regs), fields=sum(len(r["fields"]) for r in regs))
    for k, v in EXPECTED_COUNTS.items():
        if got[k] != v:
            fail("count %s = %d, expected %d" % (k, got[k], v))
    return got


def register_summary(r):
    masks = {a: 0 for a in ACCESS_TYPES}
    reset = 0
    for f in r["fields"]:
        m = field_mask(f)
        masks[f["access"]] |= m
        reset |= (f["reset"] << f["lsb"]) & m
    return masks, reset


def c_ident(text):
    return re.sub(r"[^A-Za-z0-9]", "_", text).upper()


def gen_header(regs, src):
    out = []
    w = out.append
    w("/* SPDX-License-Identifier: Apache-2.0 */")
    w("/*")
    w(" * FX1 ISP register map -- GENERATED FILE, DO NOT EDIT.")
    w(" * Generator : tools/gen_csr.py v%s" % GENERATOR_VERSION)
    w(" * Source    : %s" % src["file"])
    w(" * SHA-256   : %s" % src["sha256"])
    w(" * Map stamp : %s" % src["generated_at"])
    w(" *")
    w(" * Offsets are byte offsets relative to the ISP base address, which the")
    w(" * platform assigns. All registers are 32-bit; address bits [1:0] are ignored.")
    w(" * Deviations from the literal spreadsheet are listed in")
    w(" * docs/csr/fx1_isp_csr_schema.json (\"overrides\").")
    w(" */")
    w("#ifndef FX1_ISP_CSR_H")
    w("#define FX1_ISP_CSR_H")
    w("")
    w("#define FX1_ISP_CSR_APERTURE_BYTES 0x%05Xu" % APERTURE_BYTES)
    w("#define FX1_ISP_CSR_NUM_REGISTERS %du" % len(regs))
    w("")
    for r in regs:
        rn = "FX1_ISP_" + c_ident(r["name"])
        masks, reset = register_summary(r)
        # No spreadsheet prose here: the source is confidential (see the module
        # docstring); the worksheet row is the pointer back to the description.
        w("/* 0x%04X %s (CSR row %d) */" % (r["offset"], r["name"], r["row"]))
        w("#define %s_OFFSET 0x%04Xu" % (rn, r["offset"]))
        w("#define %s_RESET 0x%08Xu" % (rn, reset))
        for f in r["fields"]:
            fn = rn + "_" + c_ident(f["name"])
            note = "" if f["implemented"] else " (no hardware behind this field)"
            w("/*   %s [%d:%d] %s%s */" % (f["name"], f["msb"], f["lsb"], f["access"], note))
            w("#define %s_SHIFT %du" % (fn, f["lsb"]))
            w("#define %s_WIDTH %du" % (fn, f["width"]))
            w("#define %s_MASK 0x%08Xu" % (fn, field_mask(f)))
            w("#define %s_RESET 0x%Xu" % (fn, f["reset"]))
            for val, name in sorted(f["enums"].items()):
                w("#define %s_%s 0x%Xu" % (fn, c_ident(name), val))
            for bit, name in sorted(f["bits"].items()):
                if bit not in f["tied_zero_bits"]:
                    w("#define %s_%s_BIT (1u << %du)" % (fn, c_ident(name), bit))
            for al in COMPAT_ALIASES:
                if al["register"] == r["name"] and al["field"] == f["name"]:
                    kind = "_BIT" if al["new"] in f["bits"].values() else ""
                    w("#define %s_%s%s %s_%s%s /* deprecated alias (%s) */" % (
                        fn, c_ident(al["old"]), kind, fn, c_ident(al["new"]), kind, al["id"]))
        w("")
    w("#endif /* FX1_ISP_CSR_H */")
    return "\n".join(out) + "\n"


def c_str(text):
    return '"' + text.replace("\\", "\\\\").replace('"', '\\"') + '"'


def gen_table(regs, src):
    out = []
    w = out.append
    w("// SPDX-License-Identifier: Apache-2.0")
    w("// FX1 ISP register table -- GENERATED FILE, DO NOT EDIT.")
    w("// Generator: tools/gen_csr.py v%s, source SHA-256 %s" % (GENERATOR_VERSION, src["sha256"]))
    w("")
    w('#include "registers/csr_desc.h"')
    w("")
    w("namespace cdc::components::fx1_isp::csr {")
    w("namespace {")
    w("")
    for i, r in enumerate(regs):
        w("constexpr field_desc fields_%d[] = {  // %s" % (i, r["name"]))
        for f in r["fields"]:
            w("   {%s, %d, %d, 0x%08Xu, access::%s, 0x%Xu, %s, %d}," % (
                c_str(f["name"]), f["lsb"], f["width"], field_mask(f), f["access"].lower(),
                f["reset"], "true" if f["implemented"] else "false", f["row"]))
        w("};")
    w("")
    w("}  // namespace")
    w("")
    w("const reg_desc registers[] = {")
    for i, r in enumerate(regs):
        masks, reset = register_summary(r)
        w("   {0x%04Xu, %s, 0x%08Xu, 0x%08Xu, 0x%08Xu, 0x%08Xu, 0x%08Xu, 0x%08Xu, 0x%08Xu, fields_%d, %d, %d}," % (
            r["offset"], c_str(r["name"]), reset, masks["RW"], masks["RO"], masks["W1C"],
            masks["W1S"], masks["W1SC"], masks["W1P"], i, len(r["fields"]), r["row"]))
    w("};")
    w("")
    w("const std::size_t num_registers = sizeof(registers) / sizeof(registers[0]);")
    w("")
    w("}  // namespace cdc::components::fx1_isp::csr")
    return "\n".join(out) + "\n"


def gen_schema(regs, src, applied, counts):
    doc = {
        "generator": "tools/gen_csr.py", "generator_version": GENERATOR_VERSION,
        "source": src, "counts": counts, "overrides": applied, "compat_aliases": COMPAT_ALIASES,
        "access_semantics": "MAS v1.07 Table 8-178 / Table 8-180; see docs/ISP_CSR_CONTRACT.md",
        "registers": [],
    }
    for r in regs:
        masks, reset = register_summary(r)
        doc["registers"].append({
            "offset": "0x%04X" % r["offset"], "name": r["name"], "row": r["row"],
            "reset": "0x%08X" % reset,
            "masks": {a: "0x%08X" % m for a, m in masks.items()},
            "fields": [{
                "name": f["name"], "bits": "[%d:%d]" % (f["msb"], f["lsb"]),
                "access": f["access"], "reset": "0x%X" % f["reset"], "row": f["row"],
                "implemented": f["implemented"], "tied_zero_bits": f["tied_zero_bits"],
                "enums": {str(k): v for k, v in sorted(f["enums"].items())},
                "bit_names": {str(k): v for k, v in sorted(f["bits"].items())},
            } for f in r["fields"]],
        })
    return json.dumps(doc, indent=2, ensure_ascii=False) + "\n"


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--xlsx", required=True, help="path to FTEL_IP_ISP_CSR_v1.07.xlsm")
    ap.add_argument("--out", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."),
                    help="component root (default: parent of tools/)")
    ap.add_argument("--check", action="store_true",
                    help="verify the files in the tree match a fresh generation")
    ap.add_argument("--allow-new-revision", action="store_true",
                    help="accept a spreadsheet whose SHA-256 differs from the pinned one")
    args = ap.parse_args()

    digest = sha256_of(args.xlsx)
    if digest != EXPECTED_SHA256 and not args.allow_new_revision:
        fail("spreadsheet SHA-256 %s differs from pinned %s; review the new revision and "
             "rerun with --allow-new-revision" % (digest, EXPECTED_SHA256))

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    summary = {str(r[0]).strip(): r[1] for r in wb["Summary"].iter_rows(values_only=True)
               if r and r[0] is not None}
    regs = read_sheet(wb["Registers"])
    if normalise(regs) != normalise(read_sheet(wb["Flat Registers"])):
        fail("sheets 'Registers' and 'Flat Registers' disagree")
    applied = apply_overrides(regs)
    counts = validate(regs)
    src = {"file": os.path.basename(args.xlsx), "sha256": digest, "sheet": "Registers",
           "generated_at": str(summary.get("Generated At", "")),
           "module": str(summary.get("Module Name", ""))}

    outputs = {
        os.path.join("docs", "csr", "fx1_isp_csr_schema.json"): gen_schema(regs, src, applied, counts),
        os.path.join("include", "fx1_isp", "fx1_isp_csr.h"): gen_header(regs, src),
        os.path.join("src", "registers", "csr_table_gen.cpp"): gen_table(regs, src),
    }
    stale = []
    for rel, text in outputs.items():
        path = os.path.join(args.out, rel)
        if args.check:
            try:
                with open(path, encoding="utf-8") as f:
                    if f.read() != text:
                        stale.append(rel)
            except FileNotFoundError:
                stale.append(rel)
            continue
        os.makedirs(os.path.dirname(path), exist_ok=True)
        fd, tmp = tempfile.mkstemp(dir=os.path.dirname(path))
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            f.write(text)
        os.chmod(tmp, 0o644)
        os.replace(tmp, path)
        print("wrote", rel)
    if stale:
        fail("generated files out of date: " + ", ".join(stale))
    print("registers=%(registers)d fields=%(fields)d RW=%(RW)d RO=%(RO)d W1C=%(W1C)d "
          "W1S=%(W1S)d W1SC=%(W1SC)d W1P=%(W1P)d" % counts)


if __name__ == "__main__":
    main()
